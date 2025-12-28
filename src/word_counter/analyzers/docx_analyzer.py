"""
word_counter.analyzers.docx_analyzer
-----------------------------------

Analyzer for Microsoft Word (.docx) files using python-docx.

Responsibilities:
- Extract text from a .docx file (paragraphs + tables).
- Compute basic statistics:
  - number of characters
  - number of words
  - number of lines
- Extract top N most common words.

Design notes:
- Output format matches txt_analyzer.analyze_txt() so CLI/exporters can stay generic.
- We keep the logic beginner-friendly and well-documented.
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Tuple, Union

from docx import Document  # provided by python-docx


# Word token regex (same philosophy as txt_analyzer):
# - letters/numbers
# - allow apostrophes inside words (don't, user's)
_WORD_RE = re.compile(r"[A-Za-z0-9]+(?:'[A-Za-z0-9]+)?", re.UNICODE)


@dataclass(frozen=True)
class DocxAnalysisResult:
    """
    Structured result of a DOCX analysis.

    Attributes
    ----------
    file_path:
        Absolute path that was analyzed.
    file_type:
        Always "docx" for this analyzer.
    characters:
        Total number of characters in extracted text (including whitespace/newlines).
    lines:
        Total number of "lines" after extraction (we treat each extracted text line as a line).
    words:
        Total number of word tokens found.
    top_words:
        Most common words as (word, count) pairs (lowercased).
    """
    file_path: str
    file_type: str
    characters: int
    lines: int
    words: int
    top_words: List[Tuple[str, int]]

    def to_dict(self) -> Dict[str, object]:
        """Convert result to a JSON-friendly dict."""
        return {
            "file_path": self.file_path,
            "file_type": self.file_type,
            "characters": self.characters,
            "lines": self.lines,
            "words": self.words,
            "top_words": self.top_words,
        }


def _validate_docx_path(path: Union[str, Path]) -> Path:
    """
    Validate and normalize the input path.

    Raises
    ------
    FileNotFoundError
        If the file does not exist.
    IsADirectoryError
        If the path points to a directory.
    ValueError
        If the extension is not .docx (basic guard).
    """
    p = Path(path).expanduser()

    if not p.exists():
        raise FileNotFoundError(f"File not found: {p}")

    if p.is_dir():
        raise IsADirectoryError(f"Expected a file but got directory: {p}")

    if p.suffix.lower() != ".docx":
        raise ValueError(f"Expected a .docx file but got: {p.name}")

    return p.resolve()


def extract_docx_text(path: Union[str, Path]) -> str:
    """
    Extract text from a DOCX file.

    We extract:
    - Paragraph text
    - Table cell text (tables are very common in business documents)

    Parameters
    ----------
    path:
        Path to a .docx file.

    Returns
    -------
    str
        Extracted text joined by newlines.

    Notes
    -----
    - DOCX can contain many non-text elements (images, shapes). We ignore those.
    - Blank lines are removed to keep analysis clean.
    """
    p = _validate_docx_path(path)
    doc = Document(str(p))

    lines: List[str] = []

    # 1) Paragraphs
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            lines.append(text)

    # 2) Tables (rows -> cells)
    # Collect each non-empty cell as a separate line.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = (cell.text or "").strip()
                if cell_text:
                    # cell.text can contain multiple lines; normalize it
                    for part in cell_text.splitlines():
                        part = part.strip()
                        if part:
                            lines.append(part)

    return "\n".join(lines)


def tokenize_words(text: str) -> List[str]:
    """
    Convert extracted text into word tokens.

    Steps:
    - Find word-like chunks using regex
    - Lowercase them for reliable counting

    Parameters
    ----------
    text:
        Extracted text from the document.

    Returns
    -------
    list[str]
        Lowercased word tokens.
    """
    return [m.group(0).lower() for m in _WORD_RE.finditer(text)]


def analyze_docx(path: Union[str, Path], top_n: int = 10) -> Dict[str, object]:
    """
    Analyze a DOCX file and return a JSON-friendly dictionary.

    Parameters
    ----------
    path:
        Path to a .docx file.
    top_n:
        Number of most common words to return.

    Returns
    -------
    dict
        Example:
        {
          "file_path": "...",
          "file_type": "docx",
          "characters": 1234,
          "lines": 56,
          "words": 234,
          "top_words": [("the", 20), ("and", 18), ...]
        }

    Notes
    -----
    - "characters" includes whitespace and newlines (true extracted-text length).
    - "lines" is based on extracted lines (paragraphs + table cells), not Word's visual layout.
    """
    p = _validate_docx_path(path)
    text = extract_docx_text(p)

    characters = len(text)
    lines = len(text.splitlines())

    tokens = tokenize_words(text)
    words = len(tokens)

    counts = Counter(tokens)
    top_words = counts.most_common(max(0, top_n))

    result = DocxAnalysisResult(
        file_path=str(p),
        file_type="docx",
        characters=characters,
        lines=lines,
        words=words,
        top_words=top_words,
    )
    return result.to_dict()
